from django.shortcuts import render
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.db.models import Count, Q, Avg
from django.utils import timezone
from datetime import datetime, timedelta
import json

from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.platypus import Table, TableStyle, Paragraph
from reportlab.lib.styles import getSampleStyleSheet
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
import io

from control.models import Attendance, Grade
from classes.models import Lesson
from users.models import User
from administration.models import Class as SchoolClass, Subject  # Добавлен Subject

@login_required
def attendance_report(request):
    """Отчет по посещаемости"""
    # Получаем параметры фильтрации
    class_id = request.GET.get('class_id')
    start_date = request.GET.get('start_date')
    end_date = request.GET.get('end_date')
    student_id = request.GET.get('student_id')
    
    # Базовый запрос
    attendances = Attendance.objects.select_related(
        'lesson', 'lesson__class_group', 'student', 'lesson__subject'
    )
    
    # Применяем фильтры
    if class_id:
        attendances = attendances.filter(lesson__class_group_id=class_id)
    if start_date:
        attendances = attendances.filter(lesson__date__gte=start_date)
    if end_date:
        attendances = attendances.filter(lesson__date__lte=end_date)
    if student_id:
        attendances = attendances.filter(student_id=student_id)
    
    # Для учителя показываем только его уроки
    if request.user.is_teacher():
        attendances = attendances.filter(lesson__teacher=request.user)
    
    # Получаем данные для статистики
    total_attendances = attendances.count()
    present_count = attendances.filter(status='present').count()
    absent_count = attendances.filter(status='absent').count()
    late_count = attendances.filter(status='late').count()
    excused_count = attendances.filter(status='excused').count()
    
    # Группировка по статусам
    status_stats = [
        {'status': 'Присутствовал', 'count': present_count, 'color': '#28a745'},
        {'status': 'Отсутствовал', 'count': absent_count, 'color': '#dc3545'},
        {'status': 'Опоздал', 'count': late_count, 'color': '#ffc107'},
        {'status': 'Уважительная причина', 'count': excused_count, 'color': '#17a2b8'},
    ]
    
    # Группировка по дням для графика
    daily_stats = []
    if start_date and end_date:
        current_date = datetime.strptime(start_date, '%Y-%m-%d').date()
        end_date_obj = datetime.strptime(end_date, '%Y-%m-%d').date()
        
        while current_date <= end_date_obj:
            day_attendances = attendances.filter(lesson__date=current_date)
            daily_stats.append({
                'date': current_date.strftime('%d.%m'),
                'present': day_attendances.filter(status='present').count(),
                'absent': day_attendances.filter(status='absent').count(),
            })
            current_date += timedelta(days=1)
    
    # Получаем списки для фильтров
    classes = SchoolClass.objects.all()
    students = User.objects.filter(role='student')
    
    context = {
        'attendances': attendances.order_by('-lesson__date', '-lesson__start_time')[:100],
        'total_attendances': total_attendances,
        'present_count': present_count,
        'absent_count': absent_count,
        'late_count': late_count,
        'excused_count': excused_count,
        'status_stats': status_stats,
        'daily_stats': daily_stats,
        'classes': classes,
        'students': students,
        'selected_class': class_id,
        'start_date': start_date,
        'end_date': end_date,
        'selected_student': student_id,
    }
    
    return render(request, 'reports/attendance_report.html', context)

@login_required
def progress_report(request):
    """Отчет по успеваемости"""
    class_id = request.GET.get('class_id')
    subject_id = request.GET.get('subject_id')
    student_id = request.GET.get('student_id')
    period = request.GET.get('period', 'month')
    
    # Определяем период
    end_date = timezone.now().date()
    if period == 'week':
        start_date = end_date - timedelta(days=7)
    elif period == 'month':
        start_date = end_date - timedelta(days=30)
    elif period == 'quarter':
        start_date = end_date - timedelta(days=90)
    elif period == 'year':
        start_date = end_date - timedelta(days=365)
    else:
        start_date = None
    
    # Базовый запрос
    grades = Grade.objects.select_related('student', 'lesson', 'lesson__subject')
    
    # Применяем фильтры
    if class_id:
        grades = grades.filter(student__student_classes__class_group_id=class_id)
    if subject_id:
        grades = grades.filter(lesson__subject_id=subject_id)
    if student_id:
        grades = grades.filter(student_id=student_id)
    if start_date:
        grades = grades.filter(date__gte=start_date)
    
    # Для учителя показываем только его оценки
    if request.user.is_teacher():
        grades = grades.filter(teacher=request.user)
    
    # Статистика по оценкам
    total_grades = grades.count()
    avg_grade = grades.aggregate(Avg('grade'))['grade__avg']
    
    grade_counts = {
        '5': grades.filter(grade=5).count(),
        '4': grades.filter(grade=4).count(),
        '3': grades.filter(grade=3).count(),
        '2': grades.filter(grade=2).count(),
    }
    
    # Успеваемость по предметам
    subject_stats = []
    for grade in grades.values('lesson__subject__name').annotate(
        avg=Avg('grade'),
        count=Count('id')
    ).order_by('-avg'):
        subject_stats.append({
            'name': grade['lesson__subject__name'],
            'avg': round(grade['avg'], 2) if grade['avg'] else 0,
            'count': grade['count']
        })
    
    # Успеваемость по ученикам
    student_stats = []
    for grade in grades.values('student__id', 'student__last_name', 'student__first_name').annotate(
        avg=Avg('grade'),
        count=Count('id')
    ).order_by('-avg')[:20]:
        student_stats.append({
            'id': grade['student__id'],
            'name': f"{grade['student__last_name']} {grade['student__first_name']}",
            'avg': round(grade['avg'], 2) if grade['avg'] else 0,
            'count': grade['count']
        })
    
    # Получаем списки для фильтров
    classes = SchoolClass.objects.all()
    subjects = Subject.objects.all()
    students = User.objects.filter(role='student')
    
    context = {
        'total_grades': total_grades,
        'avg_grade': round(avg_grade, 2) if avg_grade else 0,
        'grade_counts': grade_counts,
        'subject_stats': subject_stats,
        'student_stats': student_stats,
        'grades': grades.order_by('-date')[:50],
        'classes': classes,
        'subjects': subjects,
        'students': students,
        'selected_class': class_id,
        'selected_subject': subject_id,
        'selected_student': student_id,
        'selected_period': period,
    }
    
    return render(request, 'reports/progress_report.html', context)

@login_required
def export_attendance_excel(request):
    """Экспорт отчета по посещаемости в Excel"""
    # Получаем параметры фильтрации как в attendance_report
    class_id = request.GET.get('class_id')
    start_date = request.GET.get('start_date')
    end_date = request.GET.get('end_date')
    
    attendances = Attendance.objects.select_related(
        'lesson', 'lesson__class_group', 'student', 'lesson__subject'
    )
    
    if class_id:
        attendances = attendances.filter(lesson__class_group_id=class_id)
    if start_date:
        attendances = attendances.filter(lesson__date__gte=start_date)
    if end_date:
        attendances = attendances.filter(lesson__date__lte=end_date)
    
    if request.user.is_teacher():
        attendances = attendances.filter(lesson__teacher=request.user)
    
    # Создаем Excel файл
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="attendance_report.xlsx"'
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Посещаемость"
    
    # Заголовки
    headers = ['Дата', 'Время', 'Класс', 'Предмет', 'Ученик', 'Статус', 'Отметил']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, size=12)
        cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal='center')
    
    # Данные
    status_translation = {
        'present': 'Присутствовал',
        'absent': 'Отсутствовал',
        'late': 'Опоздал',
        'excused': 'Уважительная причина'
    }
    
    for row, attendance in enumerate(attendances, 2):
        ws.cell(row=row, column=1, value=attendance.lesson.date.strftime('%d.%m.%Y'))
        ws.cell(row=row, column=2, value=f"{attendance.lesson.start_time.strftime('%H:%M')}-{attendance.lesson.end_time.strftime('%H:%M')}")
        ws.cell(row=row, column=3, value=str(attendance.lesson.class_group))
        ws.cell(row=row, column=4, value=attendance.lesson.subject.name)
        ws.cell(row=row, column=5, value=attendance.student.get_full_name())
        ws.cell(row=row, column=6, value=status_translation.get(attendance.status, attendance.status))
        ws.cell(row=row, column=7, value=attendance.marked_by.get_full_name() if attendance.marked_by else '')
    
    # Настройка ширины колонок
    for col in range(1, 8):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 15
    
    wb.save(response)
    return response

@login_required
def export_attendance_pdf(request):
    """Экспорт отчета по посещаемости в PDF"""
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="attendance_report.pdf"'
    
    # Создаем PDF
    p = canvas.Canvas(response, pagesize=landscape(A4))
    width, height = landscape(A4)
    
    # Заголовок
    p.setFont("Helvetica-Bold", 16)
    p.drawString(30, height - 30, "Отчет по посещаемости")
    
    # Дата
    p.setFont("Helvetica", 10)
    p.drawString(30, height - 50, f"Сформировано: {datetime.now().strftime('%d.%m.%Y %H:%M')}")
    
    # Получаем данные
    attendances = Attendance.objects.select_related(
        'lesson', 'lesson__class_group', 'student', 'lesson__subject'
    )[:50]
    
    # Создаем таблицу
    data = [['Дата', 'Класс', 'Предмет', 'Ученик', 'Статус']]
    
    status_translation = {
        'present': 'Присутствовал',
        'absent': 'Отсутствовал',
        'late': 'Опоздал',
        'excused': 'Уважительная причина'
    }
    
    for attendance in attendances:
        data.append([
            attendance.lesson.date.strftime('%d.%m.%Y'),
            str(attendance.lesson.class_group),
            attendance.lesson.subject.name,
            attendance.student.get_full_name(),
            status_translation.get(attendance.status, attendance.status)
        ])
    
    # Стили таблицы
    table = Table(data)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 12),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
        ('GRID', (0, 0), (-1, -1), 1, colors.black)
    ]))
    
    # Рисуем таблицу
    table.wrapOn(p, width - 60, height - 100)
    table.drawOn(p, 30, height - 200)
    
    p.showPage()
    p.save()
    
    return response

@login_required
def export_progress_excel(request):
    """Экспорт отчета по успеваемости в Excel"""
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="progress_report.xlsx"'
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Успеваемость"
    
    # Заголовки
    headers = ['Дата', 'Ученик', 'Класс', 'Предмет', 'Оценка', 'Тип', 'Учитель']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, size=12)
        cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal='center')
    
    # Данные
    grades = Grade.objects.select_related(
        'student', 'lesson', 'lesson__subject', 'teacher'
    )[:100]
    
    grade_type_translation = {
        'current': 'Текущая',
        'quarter': 'Четвертная',
        'final': 'Итоговая'
    }
    
    for row, grade in enumerate(grades, 2):
        ws.cell(row=row, column=1, value=grade.date.strftime('%d.%m.%Y'))
        ws.cell(row=row, column=2, value=grade.student.get_full_name())
        ws.cell(row=row, column=3, value=str(grade.lesson.class_group) if grade.lesson else '')
        ws.cell(row=row, column=4, value=grade.subject)
        ws.cell(row=row, column=5, value=grade.grade)
        ws.cell(row=row, column=6, value=grade_type_translation.get(grade.grade_type, grade.grade_type))
        ws.cell(row=row, column=7, value=grade.teacher.get_full_name() if grade.teacher else '')
    
    # Настройка ширины колонок
    for col in range(1, 8):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 15
    
    wb.save(response)
    return response
@login_required
def export_progress_pdf(request):
    """Экспорт отчета по успеваемости в PDF"""
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="progress_report.pdf"'
    
    # Создаем PDF
    p = canvas.Canvas(response, pagesize=landscape(A4))
    width, height = landscape(A4)
    
    # Заголовок
    p.setFont("Helvetica-Bold", 16)
    p.drawString(30, height - 30, "Отчет по успеваемости")
    
    # Дата
    p.setFont("Helvetica", 10)
    p.drawString(30, height - 50, f"Сформировано: {datetime.now().strftime('%d.%m.%Y %H:%M')}")
    
    # Получаем данные
    grades = Grade.objects.select_related(
        'student', 'lesson', 'lesson__subject', 'teacher'
    )[:50]
    
    # Создаем таблицу
    data = [['Дата', 'Ученик', 'Предмет', 'Оценка', 'Тип', 'Учитель']]
    
    grade_type_translation = {
        'current': 'Текущая',
        'quarter': 'Четвертная',
        'final': 'Итоговая'
    }
    
    for grade in grades:
        data.append([
            grade.date.strftime('%d.%m.%Y'),
            grade.student.get_full_name(),
            grade.subject,
            str(grade.grade),
            grade_type_translation.get(grade.grade_type, grade.grade_type),
            grade.teacher.get_full_name() if grade.teacher else ''
        ])
    
    # Стили таблицы
    table = Table(data)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 12),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
        ('GRID', (0, 0), (-1, -1), 1, colors.black)
    ]))
    
    # Рисуем таблицу
    table.wrapOn(p, width - 60, height - 100)
    table.drawOn(p, 30, height - 200)
    
    p.showPage()
    p.save()
    
    return response